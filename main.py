from typing import TypedDict, List, Dict, Any, Literal
from langgraph.graph import StateGraph, START, END
from langchain_core.messages import SystemMessage, HumanMessage
from langchain_google_genai import ChatGoogleGenerativeAI
from dotenv import load_dotenv
from pydantic import BaseModel, Field
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from io import StringIO, BytesIO
import requests
import os
import json
import streamlit as st
import time

load_dotenv()

MOCK_DATA_PATH = os.path.join(os.path.dirname(__file__), "data", "mock-sector.json")

                    

class PortfolioState(TypedDict):
    holdings: List[Dict[str, Any]] 
    market_data: Dict[str, Any]
    concentration_risk: str
    peer_analysis: str
    final_recommendation: str
    structured_recommendations: List[Dict[str, Any]]
    
class StockAction(BaseModel):
    ticker: str = Field(description="Kode ticker saham, contoh: BBCA")
    action: Literal["Buy", "Sell", "Hold"] = Field(description="Rekomendasi aksi untuk saham ini")
    alasan_singkat: str = Field(description="Alasan singkat dalam satu kalimat")

class StructuredRecommendation(BaseModel):
    recommendations: List[StockAction] = Field(description="Daftar rekomendasi aksi untuk tiap saham di portofolio")
    

llm = ChatGoogleGenerativeAI(
    model = "gemini-3.5-flash-lite",
    temperature= 0.0
)

structured_llm = ChatGoogleGenerativeAI(
    model="gemini-3.5-flash-lite",
    temperature=0.0
).with_structured_output(StructuredRecommendation)  
    
def get_company_performance(ticker: str) -> str: 
    """
        Mengambil data fundamental saham IDX untuk keperluan analisis portofolio.
        MODE=TESTING  -> baca dari mock-sector.json (hemat credit API).
        MODE=PRODUCTION -> fetch langsung dari Sectors API dan di-flatten
        ke skema yang sama persis dengan mock.
    """
    
    mode = os.getenv("MODE")
    
    if mode == "TESTING":
        with open(MOCK_DATA_PATH, "r") as f:
            mock_data = json.load(f)

        result = mock_data.get(ticker)
        
        if not result:
            return {"error": f"Data mock untuk ticker {ticker} tidak ditemukan."}
        return result
    elif mode == "PRODUCTION":
        
        api = os.getenv("SECTOR_API_KEY")
        if not api:
            return {"error": "SECTOR_API_KEY tidak ditemukan di environment variable (.env)"}
        
        url = f"https://api.sectors.app/v2/company/report/{ticker}/"
        headers = {"Authorization": api}
        
        try:
            response = requests.get(url, headers=headers)
            response.raise_for_status()
            result = response.json()
        except requests.exceptions.RequestException as e:
            return {"error": f"Gagal mengambil data untuk {ticker}: {str(e)}"}
        
        # Mengambil data (key) yang diperlukan pada sector API
        historical_list = result.get('valuation', {}).get('historical_valuation', [])
        ratio_list = result.get('financials',{}).get('historical_financial_ratio', [])
        future_list = result.get('future', {}).get('company_growth_forecasts', [])
       
       # Ambil elemen pertama secara AMAN (jika list tidak kosong)
        first_historical = historical_list[0] if historical_list else {}
        first_ratio = ratio_list[0] if ratio_list else {}
        first_future = future_list[0] if future_list else {}
        
        return {
            "symbol": result.get('symbol'),
            "company_name": result.get('company_name'),
            "sector": result.get('overview', {}).get('sector', 'N/A'),
            "pe_ratio": first_historical.get('pe', 'N/A'),
            "pbv": first_historical.get('pb', 'N/A'),
            "roe": first_ratio.get('profitability', {}).get('roe', 'N/A'),
            "revenue_growth": first_future.get('revenue_growth', 'N/A')
        }

    else:
        return {"error": "MODE tidak valid. Set MODE=TESTING atau MODE=PRODUCTION di .env"}
    

    


def data_fetcher_node(state: PortfolioState) -> PortfolioState:
    
    market_data_results = {}
 
    holdings = state["holdings"]
    
    for item in holdings:
        ticker = item["ticker"]
        data_saham = get_company_performance(ticker)
        market_data_results[ticker] = data_saham

    return {"market_data": market_data_results}


def extract_text_content(content) -> str:
    """
    Menyeragamkan response.content dari Gemini jadi string biasa.
    Kadang balik string langsung, kadang list of content blocks
    (ada metadata signature/extras yang perlu difilter).
    """
    if isinstance(content, list):
        return "".join(
            block.get("text", "") for block in content if isinstance(block, dict)
        )
    return str(content)


def risk_assessor_node(state: PortfolioState) -> PortfolioState:
    # 1. Ambil data dari state
    market_data = state.get("market_data", {})

    # 2. Siapkan prompt untuk Gemini (LLM)
    system_prompt = """Kamu adalah analis risiko portofolio saham. 
    Tugasmu: Evaluasi apakah portofolio ini terlalu terkonsentrasi pada satu sektor tertentu berdasarkan data market berikut. 
    Berikan jawaban singkat tentang risiko sektoralnya."""

    # Gabungkan prompt dengan data market_data (ubah ke string agar bisa dibaca LLM)
    pesan = [
        SystemMessage(content=system_prompt),
        HumanMessage(content=f"Data Portofolio: {json.dumps(market_data)}")
    ]


    response = llm.invoke(pesan)
    return {"concentration_risk": extract_text_content(response.content)}



def peer_analyzer_node(state: PortfolioState) -> PortfolioState:
    market_data = state.get("market_data", {})
    
    system_prompt = """Kamu adalah Analis Valuasi Saham. 
    Tugasmu: Evaluasi valuasi saham-saham berikut berdasarkan metrik PE Ratio dan PBV yang ada di data. 
    Berikan analisis singkat apakah saham tersebut tergolong undervalued (murah) atau overvalued (mahal/premium)."""

    pesan = [
        SystemMessage(content=system_prompt),
        HumanMessage(content=f"Data Fundamental: {json.dumps(market_data)}")
    ]

    response = llm.invoke(pesan)
    return {"peer_analysis": extract_text_content(response.content)}


def adviser_node(state: PortfolioState) -> PortfolioState:
  
    
    # Ambil semua data dari node sebelumnya
    holdings = state.get("holdings", [])
    risiko = state.get("concentration_risk", "")
    valuasi = state.get("peer_analysis", "")
    
    system_prompt = """You are an expert Financial Analyst AI specialized in the Indonesian Stock Market (IDX). Your primary objective is to analyze stock portfolios, evaluate financial health, and provide data-driven investment insights.

    You operate within a system where all financial data, stock prices, company fundamentals, and market news are strictly provided to you by the Sectors API. 

    CORE RULES & CONSTRAINTS:
    1. STRICT DATA RELIANCE: You must base your entire analysis ONLY on the raw data (JSON/text) provided in the user prompt. 
    2. NO HALLUCINATION: DO NOT invent, guess, or pull external financial figures, historical prices, or news that are not present in the provided data. If a metric is missing, explicitly state: "Data for [metric] is not available in the current Sectors API context."
    3. OBJECTIVITY: Provide unbiased, analytical, and professional insights. Focus on fundamental analysis, valuation, and trend identification.
    4. DISCLAIMER: Always maintain a professional boundary. Remind the user that your analysis is for informational purposes and does not constitute direct financial advice.

    WORKFLOW INSTRUCTIONS:
    When receiving a user query and the accompanying Sectors API data:
    - Step 1: Digest the provided financial metrics (e.g., Revenue, Net Income, P/E Ratio, PBV).
    - Step 2: Compare the metrics against industry standards or historical performance if the data allows.
    - Step 3: Identify key strengths (bullish factors) and weaknesses/risks (bearish factors) from the data.
    - Step 4: Synthesize a clear, actionable summary.

    OUTPUT FORMAT:
    Unless the user specifies otherwise, structure your analysis using Markdown with the following sections:
    - 📊 Executive Summary (Brief overview of the stock/portfolio health)
    - 📈 Fundamental Analysis (Breakdown of key metrics provided)
    - ⚠️ Risk Factors (Potential downsides based on the data)
    - 💡 Conclusion (Final objective thought (BUY/SELL/HOLD))

    [SYSTEM CONTEXT ENDS HERE. AWAITING USER QUERY AND SECTORS DATA]"""
    
    # Gabungkan semua konteks agar LLM bisa mengambil keputusan final
    konteks = f"""
    Portofolio Saat Ini: {json.dumps(holdings)}
    Evaluasi Risiko Sektoral: {risiko}
    Evaluasi Valuasi (Peer): {valuasi}
    """
    
    pesan = [
        SystemMessage(content=system_prompt),
        HumanMessage(content=konteks)
    ]
    
    response = llm.invoke(pesan)
    return {"final_recommendation": extract_text_content(response.content)} 

def structured_recommender_node(state: PortfolioState) -> PortfolioState:
    holdings = state.get("holdings", [])
    risiko = state.get("concentration_risk", "")
    valuasi = state.get("peer_analysis", "")

    system_prompt = """Kamu adalah Manajer Portofolio Saham.
    Berdasarkan evaluasi risiko sektoral dan valuasi yang diberikan, tentukan rekomendasi
    aksi (Beli, Jual, atau Tahan) untuk SETIAP saham dalam portofolio, beserta alasan
    singkat (1 kalimat) per saham."""

    konteks = f"""
    Portofolio Saat Ini: {json.dumps(holdings)}
    Evaluasi Risiko Sektoral: {risiko}
    Evaluasi Valuasi (Peer): {valuasi}
    """

    pesan = [
        SystemMessage(content=system_prompt),
        HumanMessage(content=konteks)
    ]

    try:
        result: StructuredRecommendation = structured_llm.invoke(pesan)
        recs = [r.model_dump() for r in result.recommendations]
    except Exception as e:
        recs = []  # fallback aman kalau LLM gagal ikutin skema

    return {"structured_recommendations": recs}        


builder = StateGraph(PortfolioState)

builder.add_node("fetcher", data_fetcher_node)
builder.add_node("risk_assessor", risk_assessor_node)
builder.add_node("peer_analyzer", peer_analyzer_node)
builder.add_node("structured_recommender", structured_recommender_node)
builder.add_node("adviser", adviser_node)


builder.add_edge(START, "fetcher")
builder.add_edge("fetcher", "risk_assessor")
builder.add_edge("risk_assessor", "peer_analyzer")
builder.add_edge("peer_analyzer", "adviser")
builder.add_edge("peer_analyzer", "structured_recommender")
builder.add_edge("adviser", END)
builder.add_edge("structured_recommender", END)

graph = builder.compile()

if __name__ == "__main__":
    intial_state = {
        "holdings": [
            {"ticker": "BBCA", "jumlah": 6000, "harga_rata": 6500},
            {"ticker": "GOTO", "jumlah": 25, "harga_rata": 50},
            {"ticker": "BBRI", "jumlah": 1000, "harga_rata": 3500},
            {"ticker": "BMRI", "jumlah": 2000, "harga_rata": 4460},
            {"ticker": "TLKM", "jumlah": 500, "harga_rata": 2500}
            
        ]
    }
    result = graph.invoke(intial_state, {"recursion_limit": 10})
    print(result.get("final_recommendation"))
    

def build_report(result: dict) -> bytes:
    holdings = result.get("holdings", [])
    market_data = result.get("market_data", {})
    structured_recs = result.get("structured_recommendations", [])
    actions_map = {rec["ticker"]: rec for rec in structured_recs}

    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio Report"

    # Judul
    ws.merge_cells("A1:I1")
    ws["A1"] = "SMART PORTFOLIO REBALANCING REPORT"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1F2933", end_color="1F2933", fill_type="solid")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    # Ringkasan risiko
    ws.merge_cells("A2:I2")
    ws["A2"] = str(result.get("concentration_risk", "-")).replace("\n", " ")
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 40

    headers = ["Ticker", "Harga Rata-rata", "Jumlah Lot", "Sector",
               "PE Ratio", "PBV", "ROE", "Rekomendasi", "Alasan"]
    header_row = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    action_colors = {"Beli": "C6F6D5", "Buy": "C6F6D5",
                      "Tahan": "FEF3C7", "Hold": "FEF3C7",
                      "Jual": "FECACA", "Sell": "FECACA"}

    row_idx = header_row + 1
    for item in holdings:
        ticker = item["ticker"]
        data = market_data.get(ticker, {})
        rec = actions_map.get(ticker, {})
        action = rec.get("action", "N/A")

        values = [ticker, item.get("harga_rata"), item.get("jumlah"),
                  data.get("sector"), data.get("pe_ratio"), data.get("pbv"),
                  data.get("roe"), action, rec.get("alasan_singkat", "-")]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if col_idx == 8 and action in action_colors:
                cell.fill = PatternFill(start_color=action_colors[action],
                                         end_color=action_colors[action], fill_type="solid")
                cell.font = Font(bold=True)
        row_idx += 1

    widths = [10, 14, 12, 14, 10, 8, 8, 12, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def get_bar_color(persen: float) -> str:
    """Tentukan warna bar berdasarkan besar alokasi."""
    if persen >= 50:
        return "#0adb3a"   # hijau
    elif persen >= 30:
        return "#d3b015"   # kuning
    else:
        return "#d20c1d"   # merah



def render_colored_bar(label: str, persen: float):
    """Render satu baris bar chart custom (ticker + persentase + bar berwarna)."""
    color = get_bar_color(persen)
    st.markdown(f"""
        <div style="margin-bottom:14px">
            <div style="display:flex;justify-content:space-between;font-size:14px;margin-bottom:4px">
                <span><b>{label}</b></span>
                <span>{persen:.0f}%</span>
            </div>
            <div style="background:#e5e5e5;border-radius:6px;height:10px;width:100%">
                <div style="background:{color};width:{min(persen, 100):.1f}%;height:10px;border-radius:6px"></div>
            </div>
        </div>
    """, unsafe_allow_html=True)


